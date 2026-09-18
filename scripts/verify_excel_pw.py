# -*- coding: utf-8 -*-
import openpyxl
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')
excel_path = sys.argv[1] if len(sys.argv) > 1 else 'test_exported_accounts.xlsx'
wb = openpyxl.load_workbook(excel_path)
missing_pw = []
checked = 0

for sheet_name in ['Công An & CAND', 'Cấp Cứu Y Tế', 'Cứu Hộ Doanh Nghiệp']:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    for r in range(6, ws.max_row + 1):
        username = str(ws.cell(row=r, column=7).value or '').strip()
        pw = str(ws.cell(row=r, column=8).value or '').strip()
        if username and not username.startswith('Tên Đăng'):
            checked += 1
            if not pw:
                missing_pw.append(f"{sheet_name} row {r}: {username}")

print(json.dumps({'checked': checked, 'missing': missing_pw}, ensure_ascii=False))
