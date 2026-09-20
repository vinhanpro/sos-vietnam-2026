# -*- coding: utf-8 -*-
"""
Generate Agency Accounts Excel (.xlsx) Report with 3 Sheets matching 2026-09-17:
Sheet 1: Công An & CAND
Sheet 2: Cấp Cứu Y Tế
Sheet 3: Cứu Hộ Doanh Nghiệp

Exactly 12 Columns:
STT, Khu Vực (Tỉnh/TP), Lực Lượng Nghiệp Vụ, Cấp Hành Chính, Tên Cơ Quan / Đơn Vị Trực Ban, Địa Bàn (Xã/Phường), Tên Đăng Nhập, Mật Khẩu, Cán Bộ Phụ Trách, Chức Vụ / Cấp Bậc, SĐT Trực Ban, SMS Tiếp Nhận
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import sys
import os

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

CITY_PROVINCES = ["TP. Hồ Chí Minh", "Hà Nội", "Đà Nẵng", "Hải Phòng", "Cần Thơ", "Huế"]

DEFAULT_PASSWORDS = {
    'admin': 'Admin',
    'congan': 'Congan@113',
    'csgt': 'Csgt@113',
    'pccc': 'Pccc@114',
    'hospital': 'Capcuu@115',
    'traffic-rescue': 'Cuuho@114',
    'police': 'Congan@113'
}

def default_password_for_account(acc):
    u = (acc.get('username') or '').lower()
    ag = (acc.get('agency') or '').lower()
    if u == 'admin': return DEFAULT_PASSWORDS['admin']
    if 'congan' in u or 'catp' in u or 'cap' in u or 'cax' in u: return DEFAULT_PASSWORDS['congan']
    if 'csgt' in u: return DEFAULT_PASSWORDS['csgt']
    if 'pccc' in u: return DEFAULT_PASSWORDS['pccc']
    if ag == 'hospital' or 'bv' in u or 'capcuu' in u: return DEFAULT_PASSWORDS['hospital']
    if ag == 'traffic-rescue' or 'cuuho' in u or 'gara' in u: return DEFAULT_PASSWORDS['traffic-rescue']
    return DEFAULT_PASSWORDS.get(ag, '2002')

def get_account_password(acc):
    pwd = acc.get('password')
    if pwd and isinstance(pwd, str) and pwd.strip() and not pwd.startswith('$') and len(pwd) < 50:
        return pwd.strip()
    init_pwd = acc.get('initialPassword')
    if init_pwd and isinstance(init_pwd, str) and init_pwd.strip():
        return init_pwd.strip()
    return default_password_for_account(acc)

def format_province_label(prov_name):
    if prov_name in CITY_PROVINCES:
        return f"TP. {prov_name.upper()}" if not prov_name.startswith("TP.") else prov_name.upper()
    return f"TỈNH {prov_name.upper()}"

def build_accounts_excel(accounts_json_path, output_excel_path):
    with open(accounts_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    if isinstance(data, dict) and 'accounts' in data and isinstance(data['accounts'], list):
        accounts = data['accounts']
        generated_at = data.get('generatedAt', '17:47:03 17/9/2026')
        officer_name = data.get('officerName', 'Trung Tâm Chỉ Huy Tác Chiến & Điều Phối Quốc Gia')
    elif isinstance(data, dict):
        accounts = list(data.values())
        generated_at = '17:47:03 17/9/2026'
        officer_name = 'Trung Tâm Chỉ Huy Tác Chiến & Điều Phối Quốc Gia'
    elif isinstance(data, list):
        accounts = data
        generated_at = '17:47:03 17/9/2026'
        officer_name = 'Trung Tâm Chỉ Huy Tác Chiến & Điều Phối Quốc Gia'
    else:
        accounts = []
        generated_at = '17:47:03 17/9/2026'
        officer_name = 'Trung Tâm Chỉ Huy Tác Chiến & Điều Phối Quốc Gia'

    wb = openpyxl.Workbook()

    # Palette & Fonts
    font_title = Font(name="Arial", size=13, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Arial", size=9.5, italic=True, color="475569")
    font_header = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
    font_group = Font(name="Arial", size=10.5, bold=True, color="0F172A")
    font_cell = Font(name="Arial", size=9, color="0F172A")
    font_mono = Font(name="Courier New", size=9.5, bold=True, color="6D28D9")

    fill_national_banner = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_police = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header_med = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    fill_header_rescue = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid")

    fill_region_group = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    fill_region_med = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_region_rescue = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # 12 CỘT CHUẨN KHỚP FILE 17/09/2026
    headers = [
        ("STT", 6),
        ("Khu Vực (Tỉnh/TP)", 18),
        ("Lực Lượng Nghiệp Vụ", 22),
        ("Cấp Hành Chính", 18),
        ("Tên Cơ Quan / Đơn Vị Trực Ban", 40),
        ("Địa Bàn (Xã/Phường)", 24),
        ("Tên Đăng Nhập", 18),
        ("Mật Khẩu", 16),
        ("Cán Bộ Phụ Trách", 22),
        ("Chức Vụ / Cấp Bậc", 18),
        ("SĐT Trực Ban", 16),
        ("SMS Tiếp Nhận", 16)
    ]

    # Categorize accounts
    police_national = []
    police_by_prov = {}
    hospital_by_prov = {}
    rescue_by_prov = {}

    for acc in accounts:
        prov = acc.get('province') or 'Cần Thơ'
        level = acc.get('level') or 'ward'
        agency = acc.get('agency') or 'police'
        username = acc.get('username') or ''

        if level == 'national' or username == 'admin':
            police_national.append(acc)
            continue

        if agency == 'hospital':
            if prov not in hospital_by_prov:
                hospital_by_prov[prov] = []
            hospital_by_prov[prov].append(acc)
        elif agency == 'traffic-rescue' or level == 'enterprise':
            if prov not in rescue_by_prov:
                rescue_by_prov[prov] = []
            rescue_by_prov[prov].append(acc)
        else: # police, csgt, fire
            if prov not in police_by_prov:
                police_by_prov[prov] = {
                    'police_prov': [],
                    'police_wards': [],
                    'csgt': [],
                    'fire': []
                }
            p = police_by_prov[prov]
            if agency == 'csgt':
                p['csgt'].append(acc)
            elif agency == 'fire':
                p['fire'].append(acc)
            else: # police
                if level == 'province' or username == 'congan' or username.startswith('catp'):
                    p['police_prov'].append(acc)
                else:
                    p['police_wards'].append(acc)

    def format_sheet_header(ws, sheet_title, sub_title, fill_hdr):
        ws.views.sheetView[0].showGridLines = True
        ws.merge_cells('A1:L1')
        ws['A1'] = "HỆ THỐNG CỨU HỘ & CẢNH BÁO SOS KHẨN CẤP QUỐC GIA (34 TỈNH THÀNH)"
        ws['A1'].font = Font(name="Arial", size=10.5, bold=True, color="FFFFFF")
        ws['A1'].fill = fill_national_banner
        ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22

        ws.merge_cells('A2:L2')
        ws['A2'] = sheet_title
        ws['A2'].font = font_title
        ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24

        ws.merge_cells('A3:L3')
        ws['A3'] = f"Thời điểm xuất: {generated_at}  ·  Đơn vị: {officer_name}  ·  {sub_title}"
        ws['A3'].font = font_subtitle
        ws['A3'].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 18

        ws.row_dimensions[4].height = 6

        header_row = 5
        ws.row_dimensions[header_row].height = 26
        for col_idx, (h_title, h_width) in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_title)
            cell.font = font_header
            cell.fill = fill_hdr
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = h_width

    # =============================================================
    # SHEET 1: Công An & CAND
    # =============================================================
    ws_police = wb.active
    ws_police.title = "Công An & CAND"
    format_sheet_header(ws_police, "DANH SÁCH TÀI KHOẢN: LỰC LƯỢNG CÔNG AN & AN NINH NHÂN DÂN", "Chế độ: BẢO MẬT TÁC CHIẾN", fill_header_police)

    cur_row = 6
    stt = 1

    def write_row(ws, acc, r_idx, prov_disp, agency_disp, level_disp, is_even, stt_val=None):
        fill = fill_zebra if is_even else fill_white
        ws.cell(row=r_idx, column=1, value=(stt_val or stt)).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=2, value=prov_disp).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=3, value=agency_disp).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=4, value=level_disp).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=5, value=acc.get('agencyName') or acc.get('unitName') or '').alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=6, value=acc.get('ward') or ('Toàn Thành Phố' if prov_disp in CITY_PROVINCES else 'Toàn Tỉnh')).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=7, value=acc.get('username', '')).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=8, value=get_account_password(acc)).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=9, value=acc.get('officerName', '')).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=10, value=acc.get('officerRank', '')).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=11, value=acc.get('officerPhone', '')).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=12, value=acc.get('officerSms', '')).alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, 13):
            c = ws.cell(row=r_idx, column=col)
            c.font = font_mono if col in [7, 8] else font_cell
            c.fill = fill
            c.border = cell_border
        ws.row_dimensions[r_idx].height = 20

    # 1.1 National Group
    if police_national:
        ws_police.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=12)
        g_cell = ws_police.cell(row=cur_row, column=1, value="🏛️ I. CẤP TRUNG ƯƠNG / QUỐC GIA (TỔNG HỢP)")
        g_cell.font = font_group
        g_cell.fill = fill_region_group
        g_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, 13): ws_police.cell(row=cur_row, column=c).border = cell_border
        ws_police.row_dimensions[cur_row].height = 22
        cur_row += 1

        for acc in police_national:
            write_row(ws_police, acc, cur_row, "Toàn Quốc", "Chỉ Huy Quốc Gia", "Trung Ương", (stt % 2 == 0))
            cur_row += 1
            stt += 1

    # 1.2 Police by Province
    r_idx = 1
    for prov_name, groups in police_by_prov.items():
        total_p = len(groups['police_prov']) + len(groups['police_wards']) + len(groups['csgt']) + len(groups['fire'])
        if total_p == 0: continue

        prov_header_text = format_province_label(prov_name)
        ws_police.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=12)
        p_cell = ws_police.cell(row=cur_row, column=1, value=f"🏙️ {r_idx + 1}. KHU VỰC: {prov_header_text}")
        p_cell.font = font_group
        p_cell.fill = fill_region_group
        p_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, 13): ws_police.cell(row=cur_row, column=c).border = cell_border
        ws_police.row_dimensions[cur_row].height = 22
        cur_row += 1
        r_idx += 1

        for acc in groups['police_prov']:
            write_row(ws_police, acc, cur_row, prov_name, "Công an khu vực", "Cấp Tỉnh/TP", (stt % 2 == 0))
            cur_row += 1
            stt += 1

        for acc in groups['police_wards']:
            write_row(ws_police, acc, cur_row, prov_name, "Công an khu vực", "Cấp Xã/Phường", (stt % 2 == 0))
            cur_row += 1
            stt += 1

        for acc in groups['csgt']:
            write_row(ws_police, acc, cur_row, prov_name, "Cảnh sát giao thông", "Cấp Phòng (PC08)", (stt % 2 == 0))
            cur_row += 1
            stt += 1

        for acc in groups['fire']:
            lvl = "Cấp Phòng (PC07)" if acc.get('level') == 'province' else "Đội PCCC Cơ Sở"
            write_row(ws_police, acc, cur_row, prov_name, "PCCC & CNCH", lvl, (stt % 2 == 0))
            cur_row += 1
            stt += 1

    # =============================================================
    # SHEET 2: Cấp Cứu Y Tế
    # =============================================================
    ws_med = wb.create_sheet(title="Cấp Cứu Y Tế")
    format_sheet_header(ws_med, "DANH SÁCH TÀI KHOẢN: HỆ THỐNG CẤP CỨU Y TẾ & BỆNH VIỆN 34 TỈNH THÀNH", "Ngành: Y TẾ ĐIỀU PHỐI", fill_header_med)

    cur_row = 6
    stt_med = 1
    m_idx = 1

    for prov_name, acc_list in hospital_by_prov.items():
        prov_header_text = format_province_label(prov_name)
        ws_med.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=12)
        m_cell = ws_med.cell(row=cur_row, column=1, value=f"🏥 {m_idx}. KHU VỰC: {prov_header_text}")
        m_cell.font = font_group
        m_cell.fill = fill_region_med
        m_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, 13): ws_med.cell(row=cur_row, column=c).border = cell_border
        ws_med.row_dimensions[cur_row].height = 22
        cur_row += 1
        m_idx += 1

        for acc in acc_list:
            lvl_code = acc.get('level') or 'province'
            lvl_disp = "Cấp Tỉnh/TP" if lvl_code == 'province' else ("Cấp Quốc Gia" if lvl_code == 'national' else "Cấp Xã/Phường/Huyện")
            write_row(ws_med, acc, cur_row, prov_name, "Cấp cứu y tế", lvl_disp, (stt_med % 2 == 0), stt_val=stt_med)
            cur_row += 1
            stt_med += 1

    # =============================================================
    # SHEET 3: Cứu Hộ Doanh Nghiệp
    # =============================================================
    ws_res = wb.create_sheet(title="Cứu Hộ Doanh Nghiệp")
    format_sheet_header(ws_res, "DANH SÁCH TÀI KHOẢN: DOANH NGHIỆP & GARAGE CỨU HỘ XE GIAO THÔNG 34 TỈNH THÀNH", "Ngành: DỊCH VỤ CỨU HỘ ĐƯỜNG BỘ", fill_header_rescue)

    cur_row = 6
    stt_res = 1
    r_idx = 1

    for prov_name, acc_list in rescue_by_prov.items():
        prov_header_text = format_province_label(prov_name)
        ws_res.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=12)
        r_cell = ws_res.cell(row=cur_row, column=1, value=f"🚗 {r_idx}. KHU VỰC: {prov_header_text}")
        r_cell.font = font_group
        r_cell.fill = fill_region_rescue
        r_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c in range(1, 13): ws_res.cell(row=cur_row, column=c).border = cell_border
        ws_res.row_dimensions[cur_row].height = 22
        cur_row += 1
        r_idx += 1

        for acc in acc_list:
            write_row(ws_res, acc, cur_row, prov_name, "Cứu hộ doanh nghiệp", "Doanh Nghiệp", (stt_res % 2 == 0), stt_val=stt_res)
            cur_row += 1
            stt_res += 1

    wb.save(output_excel_path)
    print(f"Successfully generated 3-sheet accounts excel matching 17/09 at: {output_excel_path}")

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_accounts_excel.py <json_path> <output_excel_path>")
        sys.exit(1)
    build_accounts_excel(sys.argv[1], sys.argv[2])
